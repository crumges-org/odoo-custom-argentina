import base64
import io
import json
import logging
from datetime import datetime
from odoo import models, fields, api, _
from markupsafe import Markup
from odoo.exceptions import UserError
import re

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:
    _logger.error("OpenPyXL not installed")

class L10nArImportArcaWizard(models.TransientModel):
    _name = 'l10n_ar.arca.import.wizard'
    _description = 'Asistente de Importación ARCA'

    file_data = fields.Binary(string='Archivo Excel', required=True)
    filename = fields.Char(string='Nombre Archivo')
    import_type = fields.Selection([
        ('in_invoice', 'Facturas Recibidas (Proveedores)'),
        ('out_invoice', 'Facturas Emitidas (Clientes)')
    ], string='Tipo de Importación', required=True, default='in_invoice')
    
    journal_purchase_id = fields.Many2one(
        'account.journal', 
        string='Diario de Compras',
        domain="[('type', '=', 'purchase'), ('l10n_latam_use_documents', '=', True)]"
    )
    journal_sale_id = fields.Many2one(
        'account.journal', 
        string='Diario de Ventas',
        domain="[('type', '=', 'sale'), ('l10n_latam_use_documents', '=', True)]"
    )

    line_ids = fields.One2many('l10n_ar.arca.import.line', 'wizard_id', string='Líneas')
    
    state = fields.Selection([
        ('upload', 'Carga'),
        ('review', 'Revisión'),
        ('done', 'Hecho')
    ], default='upload', string='Estado')

    company_id = fields.Many2one('res.company', string='Compañía', required=True, default=lambda self: self.env.company)
    
    pre_analysis_message = fields.Char(string='Pre-análisis', readonly=True)
    is_auto_detected = fields.Boolean(default=False)
    
    can_import = fields.Boolean(compute='_compute_can_import')

    @api.depends('line_ids.status', 'line_ids.to_import')
    def _compute_can_import(self):
        for wizard in self:
            wizard.can_import = any(l.status == 'ready' and l.to_import for l in wizard.line_ids)

    detected_period = fields.Char(string='Período Detectado', compute='_compute_detected_period')

    @api.depends('line_ids', 'line_ids.date')
    def _compute_detected_period(self):
        for wizard in self:
            dates = wizard.line_ids.mapped('date')
            if not dates:
                wizard.detected_period = False
                continue
                
            min_date = min(dates)
            max_date = max(dates)
            
            # Format dates DD/MM/AAAA
            d_from = min_date.strftime('%d/%m/%Y')
            d_to = max_date.strftime('%d/%m/%Y')
            
            # Get unique months (Spanish names if possible, but standard Babel/Python loc is safer or just specific mapping)
            # Simple approach: "Month Year"
            months = sorted(list(set(d.strftime('%m/%Y') for d in dates)))
            
            # Context-aware month names could be tricky without babel, let's use a simple mapping or Odoo's format_date if available
            # We'll stick to a simple mapping for "Mes/es detectado/s"
            month_names = {
                '01': 'Enero', '02': 'Febrero', '03': 'Marzo', '04': 'Abril', '05': 'Mayo', '06': 'Junio',
                '07': 'Julio', '08': 'Agosto', '09': 'Septiembre', '10': 'Octubre', '11': 'Noviembre', '12': 'Diciembre'
            }
            
            month_labels = []
            for m_y in months:
                m, y = m_y.split('/')
                month_labels.append(f"{month_names.get(m, m)} {y}")
            
            months_str = ", ".join(month_labels)
            
            # Determine active journal
            journal = wizard.journal_purchase_id if wizard.import_type == 'in_invoice' else wizard.journal_sale_id
            journal_name = journal.name if journal else 'No definido'
            
            pos_info = ""
            if wizard.import_type == 'out_invoice' and journal and journal.l10n_ar_afip_pos_number:
                pos_info = f" (PdV: {journal.l10n_ar_afip_pos_number})"
            
            wizard.detected_period = f"Diario: {journal_name}{pos_info} | Período detectado según facturas: {d_from} al {d_to} ({months_str})"

    # --- Dashboard Fields ---
    view_filter = fields.Selection([
        ('all', 'Todos'),
        ('ready', 'Listos'),
        ('error', 'Errores'),
        ('exists', 'Existentes')
    ], default='all', string="Filtro de Vista")
    
    analytic_account_id = fields.Many2one('account.analytic.account', string='Cuenta Analítica')

    info_all = fields.Char(compute='_compute_dashboard_data')
    info_ready = fields.Char(compute='_compute_dashboard_data')
    info_error = fields.Char(compute='_compute_dashboard_data')
    info_exists = fields.Char(compute='_compute_dashboard_data')

    visible_line_ids = fields.Many2many('l10n_ar.arca.import.line', compute='_compute_visible_lines', store=True)
    
    partner_filter_ids = fields.One2many('l10n_ar.arca.import.partner.filter', 'wizard_id', string='Filtro de Contactos')

    show_check_all = fields.Boolean(compute='_compute_mass_action_visibility')
    show_uncheck_all = fields.Boolean(compute='_compute_mass_action_visibility')

    @api.depends('line_ids.to_import', 'line_ids.status', 'view_filter', 'partner_filter_ids.selected')
    def _compute_mass_action_visibility(self):
        for wizard in self:
            # Recompute visible lines first to ensure we have the correct subset
            # Or just rely on the fact that accessing visible_line_ids triggers its compute if needed
            # But since we depend on the inputs of visible_line_ids, it should be fine.
            
            # We only consider lines that are actually 'ready' to be imported,
            # because 'error'/'exists' lines are readonly and cannot be checked.
            # The buttons apply to the subset of visible lines that are editable.
            importable_lines = wizard.visible_line_ids.filtered(lambda l: l.status == 'ready')
            
            if not importable_lines:
                wizard.show_check_all = False
                wizard.show_uncheck_all = False
            else:
                checked_count = len(importable_lines.filtered('to_import'))
                total_importable = len(importable_lines)
                
                # Show 'Mark All' if there is at least one unchecked importable line
                wizard.show_check_all = checked_count < total_importable
                
                # Show 'Unmark All' if there is at least one checked line
                wizard.show_uncheck_all = checked_count > 0

    @api.depends('line_ids', 'view_filter', 'partner_filter_ids.selected')
    def _compute_visible_lines(self):
        for wizard in self:
            domain = []
            
            # 1. Status Filter
            if wizard.view_filter == 'ready':
                domain.append(('status', '=', 'ready'))
            elif wizard.view_filter == 'error':
                domain.append(('status', '=', 'error'))
            elif wizard.view_filter == 'exists':
                domain.append(('status', '=', 'exists'))
            
            lines = wizard.line_ids.filtered_domain(domain)
            
            # 2. Partner Filter
            selected_partners = wizard.partner_filter_ids.filtered('selected')
            if selected_partners:
                # Filter lines where partner_name matches any of the selected partner names
                selected_names = selected_partners.mapped('name')
                lines = lines.filtered(lambda l: l.partner_name in selected_names)
            
            wizard.visible_line_ids = lines

    @api.onchange('partner_filter_ids', 'view_filter')
    def _onchange_filter_update(self):
        # Explicitly trigger recompute of visible lines when filters change
        # This helps in wizard context where One2many changes might not auto-propagate to other fields
        self._compute_visible_lines()

    total_lines = fields.Integer(compute='_compute_total_lines', string='Total Comprobantes en Archivo')

    @api.depends('line_ids')
    def _compute_total_lines(self):
        for wizard in self:
            wizard.total_lines = len(wizard.line_ids)

    @api.depends('line_ids', 'line_ids.status', 'partner_filter_ids.selected')
    def _compute_dashboard_data(self):
        for wizard in self:
            # Base lines: Filter by selected partners first
            # If no partner selected, use all lines
            selected_partners = wizard.partner_filter_ids.filtered('selected')
            base_lines = wizard.line_ids
            filter_label = "TODOS LOS COMPROBANTES"
            
            if selected_partners:
                # If partners are selected, filter lines
                selected_names = selected_partners.mapped('name')
                base_lines = wizard.line_ids.filtered(lambda l: l.partner_name in selected_names)
                filter_label = "COMPROBANTES SELECCIONADOS"
            
            c_all = len(base_lines)
            c_ready = len(base_lines.filtered(lambda l: l.status == 'ready'))
            c_error = len(base_lines.filtered(lambda l: l.status == 'error'))
            c_exists = len(base_lines.filtered(lambda l: l.status == 'exists'))
            
            # Format: "LABEL: COUNT"
            wizard.info_all = f"{filter_label}: {c_all}"
            wizard.info_ready = f"LISTOS PARA IMPORTAR: {c_ready}"
            wizard.info_error = f"ERRORES, NO SE IMPORTARAN: {c_error}"
            wizard.info_exists = f"EXISTENTES, YA EXISTEN: {c_exists}"

    @api.depends('line_ids', 'view_filter', 'partner_filter_ids.selected')
    def _compute_visible_lines(self):
        for wizard in self:
            # 1. Partner Filter
            selected_partners = wizard.partner_filter_ids.filtered('selected')
            base_lines = wizard.line_ids
            if selected_partners:
                selected_names = selected_partners.mapped('name')
                base_lines = wizard.line_ids.filtered(lambda l: l.partner_name in selected_names)
            
            # 2. Status Filter
            if wizard.view_filter == 'all':
                wizard.visible_line_ids = base_lines
            elif wizard.view_filter == 'ready':
                wizard.visible_line_ids = base_lines.filtered(lambda l: l.status == 'ready')
            elif wizard.view_filter == 'error':
                wizard.visible_line_ids = base_lines.filtered(lambda l: l.status == 'error')
            elif wizard.view_filter == 'exists':
                wizard.visible_line_ids = base_lines.filtered(lambda l: l.status == 'exists')
            else:
                wizard.visible_line_ids = base_lines

    def action_set_filter_all(self):
        self.view_filter = 'all'
        return self._reload_view()

    def action_set_filter_ready(self):
        self.view_filter = 'ready'
        return self._reload_view()

    def action_set_filter_error(self):
        self.view_filter = 'error'
        return self._reload_view()

    def action_set_filter_exists(self):
        self.view_filter = 'exists'
        return self._reload_view()

    def action_check_all_visible(self):
        """Mark all visible lines as to_import = True (only if ready)"""
        # We only check 'ready' lines because others are readonly/blocked
        self.visible_line_ids.filtered(lambda l: l.status == 'ready').write({'to_import': True})
        return self._reload_view()

    def action_uncheck_all_visible(self):
        """Uncheck all visible lines"""
        self.visible_line_ids.write({'to_import': False})
        return self._reload_view()

    def _reload_view(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    @api.onchange('file_data')
    def _onchange_file_data(self):
        if self.file_data:
            try:
                # Try to detect type from header
                # "Mis Comprobantes Recibidos" vs "Mis Comprobantes Emitidos"
                wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.file_data)), data_only=True)
                ws = wb.active
                # Get first cell value (A1)
                first_cell = ws.cell(row=1, column=1).value
                if first_cell and isinstance(first_cell, str):
                    detected = False
                    if 'Recibidos' in first_cell:
                        self.import_type = 'in_invoice'
                        self.pre_analysis_message = "Pre-análisis: Detectado 'Comprobantes Recibidos' (Proveedores)"
                        detected = True
                    elif 'Emitidos' in first_cell:
                        self.import_type = 'out_invoice'
                        self.pre_analysis_message = "Pre-análisis: Detectado 'Comprobantes Emitidos' (Clientes)"
                        detected = True
                    else:
                        self.pre_analysis_message = "Pre-análisis: Tipo de comprobante no detectado automáticamente"
                    
                    self.is_auto_detected = detected

                    # Trigger journal update and domain set
                    return self._onchange_import_type()
            except Exception:
                # Ignore errors here, let validation handle it
                self.pre_analysis_message = "Error en Pre-análisis: No se pudo leer el archivo"
                self.is_auto_detected = False
                pass
        else:
            self.pre_analysis_message = False
            self.is_auto_detected = False

    @api.onchange('import_type')
    def _onchange_import_type(self):
        # Clear fields
        self.journal_purchase_id = False
        self.journal_sale_id = False
        
        company = self.env.company
        
        if self.import_type == 'out_invoice':
            domain = [('type', '=', 'sale'), ('company_id', '=', company.id), ('l10n_latam_use_documents', '=', True)]
            journal = self.env['account.journal'].search(domain, limit=1)
            if journal:
                self.journal_sale_id = journal
        else:
            domain = [('type', '=', 'purchase'), ('company_id', '=', company.id), ('l10n_latam_use_documents', '=', True)]
            journal = self.env['account.journal'].search(domain, limit=1)
            if journal:
                self.journal_purchase_id = journal

    def _normalize_cuit(self, cuit):
        if not cuit:
            return False
        return str(cuit).replace('-', '').replace(' ', '')
    
    def _get_document_type(self, doc_name):
        # Mapeo básico de nombres ARCA a Document Types de Odoo
        # ARCA: "1 - Factura A", "6 - Factura B", "11 - Factura C"
        # Odoo l10n_ar: buscar por codigo
        code_match = re.match(r'^(\d+)\s-', doc_name)
        if code_match:
            code = code_match.group(1)
            # Buscar document type por codigo AFIP
            doc_type = self.env['l10n_latam.document.type'].search([
                ('code', '=', code),
                ('country_id.code', '=', 'AR')
            ], limit=1)
            return doc_type
        return False

    def _find_tax(self, amount, type_tax_use):
        # Búsqueda aproximada de impuestos por monto
        # ARCA columns: "IVA 21%", "IVA 10,5%", "IVA 27%"
        # amount es el valor porcentual (21.0, 10.5)
        # Buscar impuesto activo de ese tipo
        taxes = self.env['account.tax'].search([
            ('type_tax_use', '=', type_tax_use),
            ('amount', '=', amount),
            ('amount_type', '=', 'percent'),
            ('country_id.code', '=', 'AR'),
            ('active', '=', True)
        ])
        return taxes[0] if taxes else False

    def action_analyze(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_("Por favor suba un archivo Excel."))

        # Initialize dictionary to track unique partners for the filter tab
        unique_partners = {}

        try:
            # Decode file
            file_content = base64.b64decode(self.file_data)
        except Exception as e:
            raise UserError(_("No se pudo decodificar el archivo. Asegúrese de que sea un archivo válido. Error: %s") % e)

        # 0. Validate File Extension
        if not self.filename or not self.filename.lower().endswith('.xlsx'):
             raise UserError(_("Formato de archivo inválido. Solo se permiten archivos Excel (.xlsx)."))
            
        if 'openpyxl' not in globals():
            raise UserError(_("La librería 'openpyxl' no está instalada. Por favor contacte al administrador."))

        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.file_data)), data_only=True)
        ws = wb.active # Asumimos primera hoja
        
        # Headers mapping (based on provided sample)
        # We need to map column index to field
        headers = {}
        header_row_idx = 1
        
        rows = list(ws.iter_rows(values_only=True))
        
        start_row = 0
        
        # Expected Headers for Strict Validation based on Type
        # Common headers
        ARCA_EXPECTED_HEADERS_COMMON = [
            'Fecha', 'Tipo', 'Punto de Venta', 'Número Desde', 'Número Hasta', 'Cód. Autorización', 
            'Tipo Cambio', 'Moneda', 'Imp. Total'
        ]
        
        # Specific headers
        if self.import_type == 'in_invoice':
            # Recibidos -> Emisor info
            ARCA_EXPECTED_HEADERS_SPECIFIC = ['Tipo Doc. Emisor', 'Nro. Doc. Emisor', 'Denominación Emisor']
        else:
            # Emitidos -> Receptor info
            ARCA_EXPECTED_HEADERS_SPECIFIC = ['Tipo Doc. Receptor', 'Nro. Doc. Receptor', 'Denominación Receptor']

        ARCA_EXPECTED_HEADERS = ARCA_EXPECTED_HEADERS_COMMON + ARCA_EXPECTED_HEADERS_SPECIFIC

        found_header_row = False
        for i, row in enumerate(rows):
            if row and 'Fecha' in row and 'Tipo' in row:
                start_row = i + 1
                found_header_row = True
                
                # --- STRICT STRUCTURE VALIDATION ---
                row_values = [str(cell).strip() for cell in row if cell is not None]
                
                # Check for missing headers
                missing_headers = [h for h in ARCA_EXPECTED_HEADERS if h not in row_values]
                
                if missing_headers:
                     raise UserError(_("El archivo no tiene la estructura esperada de ARCA para '{}'.\nColumnas faltantes: {}")
                        .format(dict(self._fields['import_type'].selection).get(self.import_type), ', '.join(missing_headers)))

                # Map headers
                for col_idx, cell_val in enumerate(row):
                    if cell_val:
                        headers[cell_val] = col_idx
                break
        
        if not found_header_row:
            raise UserError(_("No se encontró la fila de encabezados válida (Fecha, Tipo, etc) o el formato no coincide con el de ARCA."))

        # --- VALIDACIÓN CUIT COMPAÑÍA ---
        # Fila 0, Celda 0 usualmente contiene: "Mis Comprobantes ... - CUIT 30718858514"
        try:
            # Re-leer fila 0 para asegurar
            first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            header_title = first_row[0]
            if header_title and 'CUIT' in str(header_title):
                # Extraer CUIT (últimos dígitos o regex)
                cuit_match = re.search(r'CUIT\s+(\d+)', str(header_title))
                if cuit_match:
                    excel_cuit = cuit_match.group(1)
                    # Normalizar Company CUIT: Remover AR, guiones y espacios
                    company_vat = self.env.company.vat or ''
                    # A veces l10n_ar agrega prefix AR or similiar
                    company_cuit = re.sub(r'\D', '', company_vat)
                    
                    if company_cuit and excel_cuit != company_cuit:
                        raise UserError(_(
                            "El CUIT del archivo ({}) no coincide con el CUIT de la compañía actual ({}).\n"
                            "Por favor verifique que está importando el archivo correcto."
                        ).format(excel_cuit, company_cuit))
        except UserError:
            raise
        except Exception as e:
            _logger.warning(f"No se pudo validar el CUIT del encabezado: {e}")
            # No bloqueamos si falla el parseo del título, pero logueamos.

        lines_values = []
        
        # Pre-fetch generic data
        company_currency = self.env.company.currency_id
        
        for row_idx in range(start_row, len(rows)):
            row = rows[row_idx]
            if not row[headers.get('Fecha')]:
                continue

            # Extraer Datos Básicos
            date_str = row[headers['Fecha']]
            # Parse Date dd/mm/yyyy
            try:
                date = datetime.strptime(date_str, '%d/%m/%Y').date() if isinstance(date_str, str) else date_str
            except:
                date = fields.Date.today()

            doc_type_name = row[headers.get('Tipo', -1)]
            pos = row[headers.get('Punto de Venta', -1)]
            number_from = row[headers.get('Número Desde', -1)]
            # number_to = row[headers.get('Número Hasta')]
            
            cuit_col = 'Nro. Doc. Emisor' if self.import_type == 'in_invoice' else 'Nro. Doc. Receptor'
            name_col = 'Denominación Emisor' if self.import_type == 'in_invoice' else 'Denominación Receptor'
            
            # Use safe get just in case logic fails or col missing despite validation
            cuit_raw = row[headers.get(cuit_col)] if headers.get(cuit_col) is not None else False
            cuit = self._normalize_cuit(cuit_raw)
            partner_name = row[headers.get(name_col)] if headers.get(name_col) is not None else 'Desconocido'
            
            # CAE Extraction
            afip_auth_code = False
            for head_alias in ['Cód. Autorización', 'CAE', 'C.A.E', 'Codigo Autorizacion', 'Código Autorización']:
                 if headers.get(head_alias) is not None:
                     afip_auth_code = row[headers[head_alias]]
                     break
            
            amount_total = row[headers.get('Imp. Total', -1)]
            
            # --- VALIDACIONES Y RESOLUCIONES ---
            status = 'ready'
            error_msgs = []
            
            # 1. Diario
            journal = False
            if self.import_type == 'in_invoice':
                journal = self.journal_purchase_id
                if not journal:
                    status = 'error'
                    error_msgs.append("Falta seleccionar Diario de Compras.")
            else:
                # Ventas:
                # User requires usage of the VALIDATED journal in the wizard, not auto-search logic.
                if self.journal_sale_id:
                    journal = self.journal_sale_id
                else:
                    status = 'error'
                    error_msgs.append(f"No se seleccionó Diario de Venta.")

            # 2. Partner (Contacto)
            # Find ALL partners with this CUIT to check for duplicates later
            partners_with_cuit = self.env['res.partner'].search([
                ('vat', '=', cuit),
                ('parent_id', '=', False)
            ])
            
            # For assignment, we prefer the first one found or create new if none
            partner = partners_with_cuit[0] if partners_with_cuit else False
            
            if not partner and not cuit:
                 status = 'error'
                 error_msgs.append("Falta CUIT para identificar contacto.")
            
            # 3. Documento Existente
            doc_type = self._get_document_type(doc_type_name)
            move = False
            
            # Ensure headers exist
            if headers.get('Número Desde') is None:
                 # Try 'Número' if 'Número Desde' missing (some variations)
                 if headers.get('Número') is not None:
                      headers['Número Desde'] = headers['Número']
                 else:
                      raise UserError("No se encuentra columna 'Número Desde' o 'Número'")

            if doc_type and journal:
                # 0. Validate Point of Sale matching for Sales (FACTURAS DE VENTA)
                if self.import_type == 'out_invoice' and journal:
                      # Check attribute existence
                      journal_pos = getattr(journal, 'l10n_ar_afip_pos_number', None)
                      
                      _logger.info(f"VALIDATION DEBUG: File POS={pos}, Journal={journal.name}, Journal POS={journal_pos}")
                      
                      if journal_pos is not None:
                          try:
                              file_pos = int(pos) if pos else 0
                              if file_pos != journal_pos:
                                  status = 'error'
                                  error_msgs.append(f"Punto de Venta incorrecto. Archivo: {file_pos}, Diario: {journal_pos}")
                          except Exception as e:
                              _logger.error(f"Error parsing/validating POS: {e}")
                      else:
                           _logger.warning("Journal has no l10n_ar_afip_pos_number set or field missing")
                
                # Robust Duplicate Check
                def _parse_int_safe(val):
                    try:
                        if isinstance(val, (int, float)):
                            return int(val)
                        if isinstance(val, str):
                            # Remove dots (thousands) and keep only digits
                            val = val.replace('.', '').replace(',', '')
                            return int(re.sub(r'\D', '', val) or 0)
                        return 0
                    except:
                        return 0

                pos_int = _parse_int_safe(pos)
                num_int = _parse_int_safe(number_from)

                try:
                    pos_int = _parse_int_safe(pos)
                    num_int = _parse_int_safe(number_from)
                except Exception as e:
                    _logger.error(f"Error parsing ints for {partner_name}: {e}")
                    pos_int = 0
                    num_int = 0

                # Formatos posibles de número de documento en Odoo (AR)
                # Standard: 00002-00001234 (5 pos, 8 number)
                # Legacy: 0002-00001234 (4 pos)
                candidates = [
                    f"{pos_int:05d}-{num_int:08d}",
                    f"{pos_int:04d}-{num_int:08d}",
                ]

                # Check against ALL partners with same CUIT if available, otherwise just general check?
                # Actually, duplicate check should be against the specific partner(s) to avoid false positives with other partners having same number?
                # But here the issue is same CUIT = same entity effectively for tax purposes.
                
                target_partners = partners_with_cuit if partners_with_cuit else (partner if partner else False)
                
                if target_partners:
                    domain = [
                        ('journal_id', '=', journal.id),
                        ('partner_id', 'in', target_partners.ids),
                        ('l10n_latam_document_type_id', '=', doc_type.id),
                        ('l10n_latam_document_number', 'in', candidates)
                    ]
                    
                    move = self.env['account.move'].search(domain, limit=1)
                
                # FALLBACK: Check without partner restriction to find duplicates assigned to WRONG partner
                if not move:
                     domain_broad = [
                        ('journal_id', '=', journal.id),
                        ('l10n_latam_document_type_id', '=', doc_type.id),
                        ('l10n_latam_document_number', 'in', candidates)
                     ]
                     move = self.env['account.move'].search(domain_broad, limit=1)
                     if move:
                         # It is a duplicate, just on another partner!
                         _logger.info(f"DUPLICATE FOUND (BROAD): {partner_name} - MATCH: {move.name} on different partner {move.partner_id.name}")
                
                # --- SAFETY CHECK ---
                # Verificar que la factura encontrada realmente coincida con los candidatos
                # Esto es para evitar "falsos positivos" si la búsqueda se comporta de forma laxa
                if move:
                    found_number = move.l10n_latam_document_number
                    if found_number not in candidates:
                        _logger.warning(f"False Positive Detection: Search found {found_number} but expected one of {candidates}. Ignoring.")
                        move = False # IGNORAR el resultado, no es duplicado real

                # Debug log
                if move:
                     _logger.info(f"DUPLICATE FOUND: {partner_name} (CUIT {cuit}) - MATCH: {move.name} ({move.l10n_latam_document_number}) ID:{move.id} on Partner: {move.partner_id.name}")
                else:
                     _logger.info(f"NO DUPLICATE: {partner_name} - Excel: {pos}-{number_from} -> Cand: {candidates}")

                if move:
                    status = 'exists'
                    # Enhanced Error Message for User Debugging
                    msg_extra = ""
                    if partner and move.partner_id != partner:
                        msg_extra = f" (En contacto: {move.partner_id.name})"
                        
                    error_msgs.append(f"<b>Factura ya existe{msg_extra}</b>: <a href='#' data-oe-model='account.move' data-oe-id='{move.id}'>{move.name}</a> (Doc: {move.l10n_latam_document_number})")

            # 4. Impuestos (Recopilar líneas)
            tax_lines = []
            
            # --- Initialize Unique Partners Dict ---
            if 'unique_partners' not in locals():
                unique_partners = {}
            # Mapeo de columnas de impuetos ARCA a tasas
            tax_mapping = {
                'IVA 21%': 21.0,
                'IVA 10,5%': 10.5,
                'IVA 27%': 27.0,
                'IVA 5%': 5.0,
                'IVA 2,5%': 2.5,
                #'IVA 0%': 0.0 # Exento or 0? 
            }
            
            # Type Tax Use verification
            type_tax_use = 'purchase' if self.import_type == 'in_invoice' else 'sale'

            invoice_lines_data = [] # Stores (tax_id, base_amount) or product line
            
            # Netos
            # ARCA has specific logic. Typically: "Neto Grav. IVA 21%", "IVA 21%"
            # We can create lines based on Netos
            
            has_taxes = False
            for col_name, rate in tax_mapping.items():
                neto_col = f"Neto Grav. {col_name}" if col_name != 'IVA 0%' else 'Neto Grav. IVA 0%' 
                # Sometimes headers change slightly "Neto Grav. IVA 21%" matches sample
                
                if headers.get(neto_col) is not None and row[headers[neto_col]]:
                    amount_neto = row[headers[neto_col]]
                    if amount_neto:
                        tax = self._find_tax(rate, type_tax_use)
                        if not tax:
                            status = 'error'
                            error_msgs.append(f"No se encontró Impuesto {rate}% {type_tax_use}.")
                        else:
                            has_taxes = True
                            invoice_lines_data.append({
                                'price_unit': amount_neto,
                                'tax_ids': [tax.id],
                                'name': f'Importe Gravado {rate}%'
                            })
            
            # Helper to find tax by name or XmlID
            def _get_tax_by_xmlid_or_name(name, type_tax_use, xmlid_suffix=False):
                tax = False
                
                # 1. Try XmlID if suffix provided (e.g. 'ri_tax_vat_no_gravado_ventas')
                if xmlid_suffix:
                    # Construct XmlID: account.{company_id}_{suffix}
                    # Example: account.5_ri_tax_vat_no_gravado_ventas
                    # Note: We need to search by external ID which is 'module.name'
                    # The module is 'account'. The name part includes company_id.
                    
                    xml_id = f"account.{self.company_id.id}_{xmlid_suffix}"
                    tax = self.env.ref(xml_id, raise_if_not_found=False)
                    if tax and tax.type_tax_use == type_tax_use:
                        return tax
                
                # 2. Fallback to Name Search
                return self.env['account.tax'].search([
                    ('name', '=', name),
                    ('type_tax_use', '=', type_tax_use),
                    ('company_id', '=', self.company_id.id)
                ], limit=1)

            # Exentos / No Gravados
            if headers.get('Neto No Gravado') is not None and row[headers['Neto No Gravado']]:
                amount_ng = row[headers['Neto No Gravado']]
                if amount_ng:
                    # Try to find specific tax for No Gravado using XmlID first
                    suffix = 'ri_tax_vat_no_gravado_compras' if type_tax_use == 'purchase' else 'ri_tax_vat_no_gravado_ventas'
                    tax_ng = _get_tax_by_xmlid_or_name('IVA No Gravado', type_tax_use, suffix)
                    
                    if not tax_ng:
                         # Fallback to search by description or similar
                         tax_ng = self.env['account.tax'].search([
                             ('name', 'ilike', 'No Grav'),
                             ('type_tax_use', '=', type_tax_use),
                             ('company_id', '=', self.company_id.id)
                         ], limit=1)
                    
                    invoice_lines_data.append({
                        'price_unit': amount_ng,
                        'tax_ids': [tax_ng.id] if tax_ng else [],
                        'name': 'Conceptos No Gravados'
                    })
            
            # Factura C / B logic (Monotributo / Consumidor Final often have no split taxes)
            # If no taxes found yet, and we have a total amount, check if we need to apply "No Corresponde" or "Exento"
            is_c_type = ' C' in doc_type_name # Simple check for Factura C, Nota de Credito C, etc.
            
            if not has_taxes and not invoice_lines_data:
                 # Check for explicit "Exento" column first
                 if headers.get('Op. Exentas') is not None and row[headers['Op. Exentas']]:
                     amount_ex = row[headers['Op. Exentas']]
                     if amount_ex:
                          suffix = 'ri_tax_vat_exento_compras' if type_tax_use == 'purchase' else 'ri_tax_vat_exento_ventas'
                          tax_ex = _get_tax_by_xmlid_or_name('IVA Exento', type_tax_use, suffix)
                          if not tax_ex:
                              tax_ex = self.env['account.tax'].search([
                                  ('name', 'ilike', 'Exento'),
                                  ('type_tax_use', '=', type_tax_use),
                                  ('company_id', '=', self.company_id.id)
                              ], limit=1)

                          invoice_lines_data.append({
                                'price_unit': amount_ex,
                                'tax_ids': [tax_ex.id] if tax_ex else [],
                                'name': 'Operaciones Exentas'
                          })
                 
                 # If still no lines and we have total, maybe it's Factura C
                 if not invoice_lines_data and amount_total:
                      if is_c_type:
                           # Apply IVA No Corresponde
                           suffix = 'ri_tax_vat_no_corresponde_compras' if type_tax_use == 'purchase' else 'ri_tax_vat_no_corresponde_ventas'
                           tax_nc = _get_tax_by_xmlid_or_name('IVA No Corresponde', type_tax_use, suffix)
                           invoice_lines_data.append({
                                'price_unit': amount_total,
                                'tax_ids': [tax_nc.id] if tax_nc else [],
                           })


            # Construct JSON for creation
            invoice_vals = {
                'ref': f"{doc_type_name} {pos}-{number_from}",
                'move_type': self.import_type,
                'invoice_date': date.strftime('%Y-%m-%d') if date else False,
                'partner_id': partner.id if partner else False, # Will need creation logic if False
                'journal_id': journal.id if journal else False,
                'l10n_latam_document_type_id': doc_type.id if doc_type else False,
                'l10n_latam_document_number': f"{int(pos):05d}-{int(number_from):08d}",
                'invoice_line_ids': [(0, 0, line) for line in invoice_lines_data]
            }
            
            if not invoice_lines_data and status == 'ready':
                 # Fallback if no specific tax columns found but total exists (Simpler inv?)
                 if amount_total:
                      invoice_vals['invoice_line_ids'] = [(0, 0, {
                          'price_unit': amount_total,
                          'tax_ids': [],
                          'name': 'Importe General'
                      })]

            # Generate Preview String (Invoice Number)
            try:
                # Clean doc type name (e.g. "001 - Factura A" -> "Factura A")
                clean_doc_name = doc_type_name
                if ' - ' in doc_type_name:
                    clean_doc_name = doc_type_name.split(' - ')[1]
                
                # Abbreviation Mapping
                abbr_map = {
                    'Factura A': 'FA-A', 'Factura B': 'FA-B', 'Factura C': 'FA-C',
                    'Nota de Crédito A': 'NC-A', 'Nota de Crédito B': 'NC-B', 'Nota de Crédito C': 'NC-C',
                    'Nota de Débito A': 'ND-A', 'Nota de Débito B': 'ND-B', 'Nota de Débito C': 'ND-C',
                    'Recibo A': 'REC-A', 'Recibo B': 'REC-B', 'Recibo C': 'REC-C',
                    'Recibo X': 'REC-X',
                }
                
                prefix = abbr_map.get(clean_doc_name, clean_doc_name)
                
                preview_str = f"{prefix} {int(pos):05d}-{int(number_from):08d}"
            except Exception as e:
                preview_str = f"{doc_type_name} {pos}-{number_from}"

            lines_values.append({
                'date': date,
                'document_type_name': doc_type_name,
                'point_of_sale': pos,
                'number': number_from,
                'cuit': cuit,
                'afip_auth_code': afip_auth_code,
                'partner_name': partner_name,
                'partner_id': partner.id if partner else False,
                'journal_id': journal.id if journal else False,
                'amount_total': amount_total,
                'currency_id': company_currency.id,
                'status': status,
                'error_desc': '<br/>'.join(error_msgs) if error_msgs else False,
                'invoice_values': json.dumps(invoice_vals),
                'move_id': move.id if move else False,
                'to_import': True if status == 'ready' else False,
                'preview_desc': preview_str
            })

            # --- Partner Filter Aggregation ---
            if partner_name not in unique_partners:
                unique_partners[partner_name] = {
                    'name': partner_name,
                    'partner_id': partner.id if partner else False,
                    'invoice_count': 0,
                    'amount_total': 0.0,
                }
            unique_partners[partner_name]['invoice_count'] += 1
            unique_partners[partner_name]['amount_total'] += amount_total

        self.line_ids = [(5, 0, 0)] + [(0, 0, val) for val in lines_values]
        
        # Populate Partner Filter
        partner_filter_values = []
        for p_name, p_data in unique_partners.items():
            partner_filter_values.append((0, 0, {
                'name': p_data['name'],
                'partner_id': p_data['partner_id'],
                'invoice_count': p_data['invoice_count'],
                'amount_total': p_data['amount_total'],
                'currency_id': company_currency.id,
            }))
        self.partner_filter_ids = [(5, 0, 0)] + partner_filter_values

        self.state = 'review'
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def action_back(self):
        self.ensure_one()
        self.state = 'upload'
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def action_import(self):
        self.ensure_one()
        created_moves = self.env['account.move']
        
        # Determine lines to process based on Partner Filter
        lines_to_process = self.line_ids
        selected_partners = self.partner_filter_ids.filtered('selected')
        if selected_partners:
            selected_names = selected_partners.mapped('name')
            lines_to_process = lines_to_process.filtered(lambda l: l.partner_name in selected_names)

        for line in lines_to_process:
            if not line.to_import or line.status != 'ready':
                continue
                
            vals = json.loads(line.invoice_values)
            
            # 1. Partner Creation if needed
            if not vals.get('partner_id'):
                # Search for CUIT identification type
                cuit_type = self.env['l10n_latam.identification.type'].search([('name', '=', 'CUIT')], limit=1)
                
                # Create partner
                partner = self.env['res.partner'].create({
                    'name': line.partner_name,
                    'vat': line.cuit,
                    'company_type': 'company', # Force Company
                    'l10n_latam_identification_type_id': cuit_type.id if cuit_type else False,
                    'l10n_ar_afip_responsibility_type_id': self.env.ref('l10n_ar.res_IVARI').id, # Default to Responsable Inscripto logic might need review but ok for now
                })
                # Chatter message for partner
                partner.message_post(body=Markup(_("Partner dado de alta automáticamente por <b>Importación de Facturas ARCA</b>.")))
                
                vals['partner_id'] = partner.id
            
            # 2. Create Move
            try:
                # If we have AFIP Auth Code, we must ensure auth_mode is CAE (especially for vendor bills or offline mode)
                if line.afip_auth_code:
                    vals['l10n_ar_afip_auth_mode'] = 'CAE'
                
                # Apply Analytic Distribution if selected
                if self.analytic_account_id and vals.get('invoice_line_ids'):
                     analytic_dist = {str(self.analytic_account_id.id): 100}
                     for cmd in vals['invoice_line_ids']:
                         if len(cmd) == 3 and isinstance(cmd[2], dict):
                             cmd[2]['analytic_distribution'] = analytic_dist

                move = self.env['account.move'].create(vals)
                
                # --- CHATTER & CAE LOGIC ---
                msg_body = Markup(_("Factura importada desde archivo ARCA: <b>%s</b>")) % (self.filename or 'Desconocido')
                move.message_post(body=msg_body)

                # Check AFIP Verification Type (if field exists and is configured)
                # Field on res.company: l10n_ar_afip_verification_type (selection)
                # 'not_available' means NO connection/validation needed. Anything else usually implies validation.
                # However, user said: "Si la opción... es distinta a 'No disponible' entonces tambien transcribe el CAE"
                
                if hasattr(self.env.company, 'l10n_ar_afip_verification_type'):
                    verif_type = self.env.company.l10n_ar_afip_verification_type
                    if verif_type != 'not_available' and line.afip_auth_code:
                        move.write({
                            'l10n_ar_afip_auth_code': line.afip_auth_code,
                            'l10n_ar_afip_result': 'A', # Always A for ARCA history
                        })
                        
                created_moves |= move
                line.status = 'exists'
                line.move_id = move.id
            except Exception as e:
                line.status = 'error'
                line.error_desc = str(e)
        
        # --- CREATE HISTORY RECORD ---
        if created_moves:
            self.env['l10n_ar.arca.import.history'].create({
                'name': f"Importación {fields.Datetime.now().strftime('%d/%m/%Y %H:%M')} - {self.filename or 'Sin Nombre'}",
                'filename': self.filename,
                'file_data': self.file_data,
                'import_type': self.import_type,
                'move_ids': [(6, 0, created_moves.ids)]
            })

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _("Importación completada"),
                'message': _(f"Se crearon {len(created_moves)} facturas."),
                'type': 'success',
                'sticky': False,
                'next': {'type': 'ir.actions.act_window_close'}
            }
        }

class L10nArImportArcaLine(models.TransientModel):
    _name = 'l10n_ar.arca.import.line'
    _description = 'Línea de Importación ARCA'
    _order = 'date desc, number desc'

    wizard_id = fields.Many2one('l10n_ar.arca.import.wizard', string='Wizard')
    
    date = fields.Date(string='Fecha')
    document_type_name = fields.Char(string='Tipo Doc.')
    point_of_sale = fields.Char(string='Pto. Venta')
    number = fields.Char(string='Número')
    cuit = fields.Char(string='CUIT')
    
    partner_name = fields.Char(string='Razón Social')
    partner_id = fields.Many2one('res.partner', string='Partner')
    
    amount_total = fields.Monetary(string='Total', currency_field='currency_id')
    currency_id = fields.Many2one('res.currency', string='Moneda')
    
    status = fields.Selection([
        ('ready', 'Listo'),
        ('exists', 'Ya Existe'),
        ('error', 'Error')
    ], string='Estado', default='ready')
    
    error_desc = fields.Html(string='Detalle Error')
    
    # Store JSON data for invoice creation
    invoice_values = fields.Text(string='Valores JSON')
    
    move_id = fields.Many2one('account.move', string='Factura Creada')
    journal_id = fields.Many2one('account.journal', string='Diario')
    
    afip_auth_code = fields.Char(string='CAE')
    
    to_import = fields.Boolean(string='Importar', default=False)
    
    preview_desc = fields.Char(string='Vista Previa')
    
    
    def action_toggle_import(self):
        for line in self:
            if line.status == 'ready':
                line.to_import = not line.to_import
        
        # Reload the wizard view to reflect changes
        # Use wizard_id to target the correct record
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'l10n_ar.arca.import.wizard',
            'res_id': self[0].wizard_id.id,
            'view_mode': 'form',
            'target': 'new',
        }
